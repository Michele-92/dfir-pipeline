"""Test fuer Punkt 1 — reboot_sessions.xlsx, Kernel-Quelle pfadgenau.

Verifiziert in der Quelle-Spalte der Reboot_X-Mappen:
  - source='kernel' erscheint als nachpruefbarer Pfad /var/log/kern.log
    (nicht als nacktes Wort 'kernel')
  - bevorzugt orig_path (Extraktionspfad), markiert Fallback mit '*'
"""
import sys
from datetime import datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("openpyxl")

from openpyxl import load_workbook

from models.pipeline_context import PipelineContext
from models.event import ForensicEvent
from utils.event_store import EventStore
from stages import stage14_export


def _ctx(tmp_path: Path) -> PipelineContext:
    db = tmp_path / 'events.db'
    evs = [
        # Boot/Shutdown -> bilden eine Session
        ForensicEvent(datetime(2023, 3, 7, 8, 0, 0), 'kernel', 'system_boot',
                      'Linux version 5.15 system boot', severity='info',
                      orig_path='/var/log/kern.log'),
        ForensicEvent(datetime(2023, 3, 7, 12, 0, 0), 'kernel', 'system_shutdown',
                      'reached target shutdown', severity='info',
                      orig_path='/var/log/kern.log'),
        # In-Session HIGH-Event mit orig_path -> echte Quelle
        ForensicEvent(datetime(2023, 3, 7, 9, 0, 0), 'kernel', 'kernel_event',
                      'segfault in process', severity='high',
                      orig_path='/var/log/kern.log'),
        # In-Session HIGH-Event OHNE orig_path -> Fallback mit '*'
        ForensicEvent(datetime(2023, 3, 7, 10, 0, 0), 'kernel', 'kernel_event',
                      'kernel oops detected', severity='high',
                      orig_path=''),
    ]
    with EventStore(db) as store:
        store.insert_events(evs)
    ctx = PipelineContext()
    ctx.events_db_path = db
    ctx.case_dir = tmp_path
    return ctx


def _quelle_spalte(xlsx: Path):
    wb = load_workbook(xlsx)
    assert 'Reboot_1' in wb.sheetnames, f'keine Reboot_1-Mappe: {wb.sheetnames}'
    ws = wb['Reboot_1']
    # Header in Zeile 2, Spalte 6 = 'Quelle (Pfad)'
    assert ws.cell(2, 6).value == 'Quelle (Pfad)'
    return [ws.cell(r, 6).value for r in range(3, ws.max_row + 1)
            if ws.cell(r, 6).value]


def test_kernel_quelle_ist_pfad(tmp_path):
    ctx = _ctx(tmp_path)
    stage14_export._write_reboot_sessions_excel(ctx, tmp_path)
    quellen = _quelle_spalte(tmp_path / 'reboot_sessions.xlsx')
    # nirgends das nackte Wort 'kernel'
    assert 'kernel' not in quellen
    # echter Pfad vorhanden
    assert '/var/log/kern.log' in quellen


def test_fallback_markiert(tmp_path):
    ctx = _ctx(tmp_path)
    stage14_export._write_reboot_sessions_excel(ctx, tmp_path)
    quellen = _quelle_spalte(tmp_path / 'reboot_sessions.xlsx')
    # Event ohne orig_path -> Standardpfad mit '*'
    assert '/var/log/kern.log *' in quellen
