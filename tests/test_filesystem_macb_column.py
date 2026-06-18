"""Test fuer Punkt 2 — filtered_filesystem_timeline.xlsx.

Verifiziert die kompakte MACB-Spalte (Timestamp-Typ m/a/c/b je Datei):
  - Datei mit allen vier Zeitstempeln  -> 'macb'
  - Datei mit nur m und c              -> 'm.c.'
  - die vier Einzelspalten mtime..btime sind vorhanden
"""
import sys
from datetime import datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("openpyxl")

from openpyxl import load_workbook

from models.pipeline_context import PipelineContext
from models.event import ForensicEvent
from utils.event_store import EventStore
from stages import stage14_export


def _mt(fp, letter, day):
    # mactime-Event: source='mactime', event_type='filesystem_<letter>'
    return ForensicEvent(datetime(2023, 3, day, 12, 0, 0), 'mactime',
                         f'filesystem_{letter}', f'{letter} change {fp}',
                         severity='info', orig_path=fp)


def _ctx(tmp_path: Path) -> PipelineContext:
    db = tmp_path / 'events.db'
    evs = []
    # /etc/machine-id mit allen vier Zeitstempeln
    for i, letter in enumerate(('m', 'a', 'c', 'b'), 1):
        e = _mt('/etc/machine-id', letter, i + 1)
        e.file_path = '/etc/machine-id'
        evs.append(e)
    # /etc/hostname nur mtime + ctime
    for letter in ('m', 'c'):
        e = _mt('/etc/hostname', letter, 7)
        e.file_path = '/etc/hostname'
        evs.append(e)
    with EventStore(db) as store:
        store.insert_events(evs)
    ctx = PipelineContext()
    ctx.events_db_path = db
    ctx.case_dir = tmp_path
    ctx.forensic_findings = []
    return ctx


def _timeline_rows(xlsx: Path):
    wb = load_workbook(xlsx)
    ws = wb['Timeline']
    assert ws.cell(2, 1).value == 'Datei_Pfad'
    assert ws.cell(2, 2).value == 'MACB'
    assert ws.cell(2, 3).value == 'mtime_UTC'
    assert ws.cell(2, 6).value == 'btime_UTC'
    out = {}
    for r in range(3, ws.max_row + 1):
        fp = ws.cell(r, 1).value
        if fp:
            out[fp] = ws.cell(r, 2).value   # MACB-Spalte
    return out


def test_macb_spalte(tmp_path):
    ctx = _ctx(tmp_path)
    stage14_export._write_filtered_filesystem_timeline_excel(ctx, tmp_path)
    rows = _timeline_rows(tmp_path / 'filtered_filesystem_timeline.xlsx')
    assert rows.get('/etc/machine-id') == 'macb'
    assert rows.get('/etc/hostname') == 'm.c.'
