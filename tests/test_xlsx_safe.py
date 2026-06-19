"""Test fuer den openpyxl-Steuerzeichen-Schutz (Stage-14-Crash-Fix).

Reproduziert den gemeldeten Fehler: ein IOC mit Binaer-/Steuerzeichen sprengte
beim Schreiben der iocs.xlsx den gesamten Export (IllegalCharacterError).
Mit utils.xlsx_safe darf das NICHT mehr passieren.
"""
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook


def test_steuerzeichen_kein_crash():
    import utils.xlsx_safe  # noqa: F401  aktiviert den Schutz
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "boese\x15\x07\x1ewert ok")        # vorher: IllegalCharacterError
    assert ws.cell(1, 1).value == "boesewert ok"     # Steuerzeichen entfernt


def test_iocs_excel_mit_binaer_ioc():
    from models.pipeline_context import PipelineContext
    from models.ioc import IOC
    from stages import stage14_export   # importiert utils.xlsx_safe mit
    case = Path(tempfile.mkdtemp())
    ctx = PipelineContext(case_dir=case)
    ctx.iocs = [
        IOC(type="url", value="http://evil\x15\x07.com/x", source="bulk_extractor",
            context="binaerer\x1e Treffer", timestamp=datetime.now(tz=timezone.utc)),
        IOC(type="ip", value="8.8.8.8", source="auth", context="dns",
            timestamp=datetime.now(tz=timezone.utc)),
    ]
    # Darf NICHT mehr abstuerzen
    stage14_export._write_iocs_excel(ctx, case)
    out = case / "iocs.xlsx"
    assert out.exists()
    wb = load_workbook(out)
    assert len(wb.sheetnames) >= 1
