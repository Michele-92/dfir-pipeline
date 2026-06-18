"""Test fuer Punkt 5 — 3 Excel-Auszuege im modernen Report.

Legt die drei Beilagen (ip_sessions / reboot_sessions /
filtered_filesystem_timeline) synthetisch an und prueft, dass deren
wichtigste Zeilen als Auszug im PDF erscheinen (inkl. BEILAGE-Verweis).
"""
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip('pypdf')
pytest.importorskip('reportlab')
pytest.importorskip('openpyxl')

from pypdf import PdfReader
from openpyxl import Workbook

from models.pipeline_context import PipelineContext
from models.event import ForensicEvent
from stages.report_modern import build_modern_report, _xlsx_rows


def _xlsx(path, sheet, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(1, 1, 'BANNER')                       # Zeile 1 = Banner
    for c, h in enumerate(header, 1):
        ws.cell(2, c, h)                          # Zeile 2 = Header
    for r, row in enumerate(rows, 3):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(str(path))


def _make_excels(case: Path):
    _xlsx(case / 'ip_sessions.xlsx', 'Externe IPs',
          ['IP', 'Typ', 'Login_Anzahl', 'Login_Methoden', 'Erster_Login_UTC',
           'Letzter_Login_UTC', 'Session_Dauer', 'Benutzer', 'Log-Quellen',
           'In_AuthLog', 'In_Journal', 'In_Wtmp', 'In_Wtmpdb'],
          [['203.0.113.7', 'Extern', 4, 'Remote', '2023-03-07 08:00',
            '2023-03-07 09:00', '1h', 'root', 'wtmpdb',
            '✗', '✗', '✗', '✓']])
    _xlsx(case / 'reboot_sessions.xlsx', 'Übersicht',
          ['#', 'Boot_Start_UTC', 'Shutdown_UTC', 'Laufzeit', 'Ereignisse', 'Hinweis'],
          [[1, '2023-03-07 08:00:00', '2023-03-07 08:20:00', '< 30 Min', 87,
            '⚠ Kurze Laufzeit — verdächtiger Neustart']])
    _xlsx(case / 'filtered_filesystem_timeline.xlsx', 'Timeline',
          ['Datei_Pfad', 'MACB', 'mtime_UTC', 'atime_UTC', 'ctime_UTC', 'btime_UTC',
           'Auffälligkeit', 'Severity', 'Kategorie'],
          [['/root/staging.sh', 'macb', '2023-03-07 09:00:00', '2023-03-07 09:01:00',
            '2023-03-07 10:00:00', '2023-03-07 11:00:00',
            '⚠ btime > mtime → kopiert?', 'HIGH', 'System']])


def _report_text():
    case = Path(tempfile.mkdtemp()) / 'case'
    case.mkdir()
    _make_excels(case)
    ctx = PipelineContext(case_dir=case, combined_case=False)
    ctx.evidence_items = [{'name': 'img.E01', 'os_name': 'Ubuntu 22.04',
                           'hostname': 'host', 'file_type': 'E01', 'file_size_gb': 10.0,
                           'partition_layout': [], 'partition_profiles': [{}]}]
    ctx.normalized_events = [
        ForensicEvent(timestamp=datetime(2023, 3, 7, 9, 0, tzinfo=timezone.utc),
                      source='auth', event_type='ssh', message='Login root',
                      severity='high', orig_path='/var/log/auth.log')]
    ctx.parser_stats = {'auth': 10}
    ctx.parsed_events = 10
    ctx.stage_status = {'quality': 'GUT'}
    build_modern_report(ctx, case)
    pdf = case / 'forensischer_analysebericht.pdf'
    assert pdf.exists(), 'PDF wurde nicht erzeugt'
    return '\n'.join(p.extract_text() for p in PdfReader(str(pdf)).pages)


def test_xlsx_rows_liest_header_und_daten(tmp_path):
    p = tmp_path / 't.xlsx'
    _xlsx(p, 'S', ['A', 'B'], [[1, 2], [3, 4]])
    header, data = _xlsx_rows(p, 'S')
    assert header == ['A', 'B']
    assert data == [[1, 2], [3, 4]]
    # fehlendes Blatt -> leer
    assert _xlsx_rows(p, 'fehlt') == ([], [])


def test_ip_auszug_im_report():
    txt = _report_text()
    assert 'Anmelde-Sitzungen' in txt
    assert '203.0.113.7' in txt
    assert 'ip_sessions.xlsx' in txt


def test_reboot_auszug_im_report():
    txt = _report_text()
    assert 'Reboot-Sitzungen' in txt
    assert 'reboot_sessions.xlsx' in txt


def test_filesystem_auszug_im_report():
    txt = _report_text()
    assert 'Dateisystem-Timeline' in txt
    assert '/root/staging.sh' in txt
    assert 'filtered_filesystem_timeline.xlsx' in txt
